# Integrity check of the built decks (python-pptx): opens each file, counts slides, embedded pictures, notes,
# conceptual-diagram labels and checks picture aspect ratios against the embedded image. Writes tools/check-report.json.
import json, glob, os, sys, io
from pptx import Presentation
from pptx.util import Emu
from PIL import Image
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '../..'))
out = {}
for f in sorted(glob.glob(os.path.join(ROOT, '[01][0-9]-*.pptx'))):
    p = Presentation(f)
    slides = len(p.slides); pics = 0; screenshots = 0; notes = 0; distorted = []; diagrams = 0; empty = []
    for i, s in enumerate(p.slides, 1):
        txt = ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if 'Conceptual diagram' in txt: diagrams += 1
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip(): notes += 1
        shapes = list(s.shapes)
        if len(shapes) < 2 and not txt.strip(): empty.append(i)
        for sh in shapes:
            if sh.shape_type == 13:  # picture
                pics += 1
                if (sh.name or '').startswith('Screenshot'): screenshots += 1
                im = Image.open(io.BytesIO(sh.image.blob))
                r_img = im.width / im.height; r_box = sh.width / sh.height
                if abs(r_img - r_box) / r_img > 0.01: distorted.append((i, sh.name, round(r_img, 3), round(r_box, 3)))
    out[os.path.basename(f)] = dict(slides=slides, pictures=pics, screenshot_placements=screenshots, slides_with_notes=notes,
        slides_without_notes=slides - notes, conceptual_diagram_slides=diagrams, distorted=distorted, empty=empty,
        size_mb=round(os.path.getsize(f) / 1e6, 1))
    print(f"{os.path.basename(f)}: {slides} slides, {screenshots} screenshot placements, notes {notes}/{slides}, diagrams {diagrams}, distorted {len(distorted)}, empty {empty}, {out[os.path.basename(f)]['size_mb']} MB")
json.dump(out, open(os.path.join(ROOT, 'tools/check-report.json'), 'w'), indent=2)
